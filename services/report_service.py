from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class ReportService:
    def __init__(self, output_dir="./data/reports"):

        self.output_dir = Path(output_dir)

        self.output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    # =========================================================
    # PUBLIC METHOD
    # =========================================================

    def create_client_report(
        self,
        state,
        session_id,
    ):

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        safe_session_id = "".join(
            ch if ch.isalnum() or ch in ("-", "_") else "_"
            for ch in str(session_id)
        ) or "session"
        filename = f"patent_analysis_{safe_session_id}_{timestamp}.docx"

        report_path = self.output_dir / filename

        self._build_document(
            state,
            session_id,
            report_path,
        )

        return str(report_path)

    # =========================================================
    # BUILD DOCUMENT
    # =========================================================

    def _build_document(
        self,
        state,
        session_id,
        report_path,
    ):

        # LangGraph state values may be dictionaries or Pydantic objects.
        invention = self._get(state, "invention")
        if invention is None:
            raise ValueError("Cannot create patent report: state is missing 'invention'.")

        keywords = self._get(state, "keywords")
        cpc_analysis = self._get(state, "cpc_analysis")
        uspto_queries = self._get(state, "uspto_queries")
        espacenet_queries = self._get(state, "espacenet_queries")
        npl_results = self._get(state, "npl_results", []) or []
        npl_summary = self._get(state, "npl_search_summary", {}) or {}

        document = Document()

        self._configure_page(document)

        self._add_header(document)

        self._add_footer(document)

        self._add_cover(
            document,
            invention,
            session_id,
        )

        document.add_page_break()

        # -----------------------------------------------------
        # 1. Executive Summary
        # -----------------------------------------------------

        self._add_executive_summary(
            document,
            invention,
        )

        # -----------------------------------------------------
        # 2. Invention Decomposition
        # -----------------------------------------------------

        self._add_invention_decomposition(
            document,
            invention,
        )

        # -----------------------------------------------------
        # 3. Search Vocabulary
        # -----------------------------------------------------

        if keywords:
            self._add_search_vocabulary(document, keywords)
        else:
            document.add_heading("3. Search Vocabulary", level=1)
            document.add_paragraph("Search vocabulary was not available in the graph state for this run.")

        # -----------------------------------------------------
        # 4. CPC Analysis
        # -----------------------------------------------------

        if cpc_analysis:
            self._add_cpc_analysis(document, cpc_analysis)
        else:
            document.add_heading("4. CPC Classification Analysis", level=1)
            document.add_paragraph("CPC analysis was not available in the graph state for this run.")

        # -----------------------------------------------------
        # 5. USPTO Queries
        # -----------------------------------------------------

        if uspto_queries:
            self._add_search_queries(
                document,
                uspto_queries,
                "5. USPTO Patent Public Search Queries",
                "USPTO",
            )

        # -----------------------------------------------------
        # 6. Espacenet Queries
        # -----------------------------------------------------

        if espacenet_queries:
            self._add_search_queries(
                document,
                espacenet_queries,
                "6. Espacenet Search Queries",
                "Espacenet",
            )

        # -----------------------------------------------------
        # 7. Non-Patent Literature (NPL) Search Results
        # -----------------------------------------------------

        self._add_npl_results(document, npl_results, npl_summary)

        # -----------------------------------------------------
        # 8. Methodology
        # -----------------------------------------------------

        self._add_methodology(document, npl_summary=npl_summary)

        # -----------------------------------------------------
        # 9. Traceability
        # -----------------------------------------------------

        self._add_traceability(document, session_id, npl_summary=npl_summary)

        # -----------------------------------------------------
        # Confidentiality
        # -----------------------------------------------------

        self._add_confidentiality_notice(
            document,
        )

        document.save(report_path)

    # =========================================================
    # PAGE CONFIGURATION
    # =========================================================

    def _configure_page(
        self,
        document,
    ):

        section = document.sections[0]

        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        document.styles["Normal"].font.name = "Aptos"
        document.styles["Normal"].font.size = Pt(10)

    # =========================================================
    # HEADER
    # =========================================================

    def _add_header(
        self,
        document,
    ):

        header = document.sections[0].header

        paragraph = header.paragraphs[0]

        paragraph.text = "GVG Consulting Services | Patent Invention Analysis"

        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for run in paragraph.runs:
            run.font.size = Pt(8)

    # =========================================================
    # FOOTER
    # =========================================================

    def _add_footer(
        self,
        document,
    ):

        footer = document.sections[0].footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run("Confidential – Prepared for Client")

        run.font.size = Pt(8)

    # =========================================================
    # COVER
    # =========================================================

    def _add_cover(
        self,
        document,
        invention,
        session_id,
    ):

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run("PATENT INVENTION ANALYSIS REPORT")

        run.bold = True
        run.font.size = Pt(24)

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run("Invention Decomposition • Search Vocabulary • CPC Classification • Patent Search Queries")

        run.italic = True
        run.font.size = Pt(11)

        document.add_paragraph()

        table = document.add_table(
            rows=4,
            cols=2,
        )

        table.style = "Table Grid"

        rows = [
            (
                "Client / Matter",
                "[CLIENT / MATTER NAME]",
            ),
            (
                "Session / Reference",
                session_id,
            ),
            (
                "Invention Title",
                self._get(invention, "title", "[INVENTION TITLE]"),
            ),
            (
                "CPC Corpus Version",
                "2026.08",
            ),
        ]

        for row, values in zip(
            table.rows,
            rows,
        ):
            self._set_cell(
                row.cells[0],
                values[0],
                bold=True,
            )

            self._shade_cell(
                row.cells[0],
                "E7E6E6",
            )

            self._set_cell(
                row.cells[1],
                values[1],
            )

        document.add_paragraph()

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run("CONFIDENTIAL CLIENT DELIVERABLE")

        run.bold = True

    # =========================================================
    # 1. EXECUTIVE SUMMARY
    # =========================================================

    def _add_executive_summary(
        self,
        document,
        invention,
    ):

        document.add_heading(
            "1. Executive Summary",
            level=1,
        )

        document.add_paragraph(str(self._get(invention, "invention_summary", "Summary not available.")))

    # =========================================================
    # 2. INVENTION DECOMPOSITION
    # =========================================================

    def _add_invention_decomposition(
        self,
        document,
        invention,
    ):

        document.add_heading(
            "2. Invention Decomposition",
            level=1,
        )

        document.add_paragraph("The invention has been decomposed into structural, procedural/process and functional features to support subsequent patent searching and classification.")

        self._add_feature_table(
            document,
            "2.1 Structural Features",
            self._get(invention, "structural_features", []) or [],
            "Structural",
        )

        self._add_feature_table(
            document,
            "2.2 Procedural / Process Features",
            self._get(invention, "procedural_features", []) or [],
            "Procedural",
        )

        self._add_feature_table(
            document,
            "2.3 Functional Features",
            self._get(invention, "functional_features", []) or [],
            "Functional",
        )

    # =========================================================
    # FEATURE TABLE
    # =========================================================

    def _add_feature_table(
        self,
        document,
        heading,
        features,
        feature_type,
    ):

        document.add_heading(
            heading,
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=3,
        )

        table.style = "Table Grid"

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = [
            "Feature ID",
            "Feature",
            "Technical Details",
        ]

        for index, header in enumerate(headers):
            self._set_cell(
                table.rows[0].cells[index],
                header,
                bold=True,
            )

            self._shade_cell(
                table.rows[0].cells[index],
                "D9EAF7",
            )

        for item in features:
            row = table.add_row()

            self._set_cell(
                row.cells[0],
                self._get(item, "feature_id", ""),
            )

            if feature_type == "Structural":
                description = self._get(item, "feature", "")

            elif feature_type == "Procedural":
                description = self._get(item, "step", "")

            elif feature_type == "Functional":
                description = self._get(item, "function", "")

            else:
                description = ""

            self._set_cell(
                row.cells[1],
                description,
            )

            self._set_cell(
                row.cells[2],
                self._get(item, "details", ""),
            )

        document.add_paragraph()

    # =========================================================
    # 3. SEARCH VOCABULARY
    # =========================================================

    def _add_search_vocabulary(
        self,
        document,
        keywords,
    ):

        document.add_heading(
            "3. Search Vocabulary",
            level=1,
        )

        document.add_paragraph("Technical search terminology was generated independently for structural, procedural/process and functional features and then consolidated.")

        self._add_keyword_table(
            document,
            "3.1 Structural Keywords",
            self._get(keywords, "structural_keywords", []) or [],
        )

        self._add_keyword_table(
            document,
            "3.2 Procedural / Process Keywords",
            self._get(keywords, "procedural_keywords", []) or [],
        )

        self._add_keyword_table(
            document,
            "3.3 Functional Keywords",
            self._get(keywords, "functional_keywords", []) or [],
        )

        document.add_heading(
            "3.4 Consolidated Search Vocabulary",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        self._set_cell(
            table.rows[0].cells[0],
            "No.",
            bold=True,
        )

        self._set_cell(
            table.rows[0].cells[1],
            "Consolidated Keyword / Term",
            bold=True,
        )

        self._shade_cell(
            table.rows[0].cells[0],
            "D9EAF7",
        )

        self._shade_cell(
            table.rows[0].cells[1],
            "D9EAF7",
        )

        for index, keyword in enumerate(
            self._get(keywords, "consolidated_keywords", []) or [],
            start=1,
        ):
            row = table.add_row()

            self._set_cell(
                row.cells[0],
                index,
            )

            self._set_cell(
                row.cells[1],
                keyword,
            )

        document.add_paragraph()

    # =========================================================
    # KEYWORD TABLE
    # =========================================================

    def _add_keyword_table(
        self,
        document,
        heading,
        keyword_sets,
    ):

        document.add_heading(
            heading,
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=4,
        )

        table.style = "Table Grid"

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = [
            "Feature ID",
            "Feature",
            "Primary Keywords",
            "Synonyms / Alternative Terms",
        ]

        for index, header in enumerate(headers):
            self._set_cell(
                table.rows[0].cells[index],
                header,
                bold=True,
            )

            self._shade_cell(
                table.rows[0].cells[index],
                "D9EAF7",
            )

        for item in keyword_sets:
            row = table.add_row()

            self._set_cell(
                row.cells[0],
                self._get(item, "feature_id", ""),
            )

            self._set_cell(
                row.cells[1],
                self._get(item, "feature", ""),
            )

            self._set_cell(
                row.cells[2],
                "; ".join(self._as_list(self._get(item, "primary_keywords", []))),
            )

            alternatives = (
                self._as_list(self._get(item, "synonyms", []))
                + self._as_list(self._get(item, "alternative_terms", []))
            )

            self._set_cell(
                row.cells[3],
                "; ".join(alternatives),
            )

        document.add_paragraph()

    # =========================================================
    # 4. CPC ANALYSIS
    # =========================================================

    def _add_cpc_analysis(
        self,
        document,
        cpc_analysis,
    ):

        document.add_heading(
            "4. CPC Classification Analysis",
            level=1,
        )

        document.add_paragraph("Candidate CPC classifications were retrieved from the CPC 2026.08 corpus using semantic similarity and subsequently evaluated for technical relevance.")

        table = document.add_table(
            rows=1,
            cols=5,
        )

        table.style = "Table Grid"

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = [
            "CPC Code",
            "Classification / Title",
            "Feature(s)",
            "Relevance",
            "Technical Rationale",
        ]

        for index, header in enumerate(headers):
            self._set_cell(
                table.rows[0].cells[index],
                header,
                bold=True,
            )

            self._shade_cell(
                table.rows[0].cells[index],
                "D9EAD3",
            )

        for candidate in (self._get(cpc_analysis, "candidates", []) or []):
            row = table.add_row()

            self._set_cell(
                row.cells[0],
                self._get(candidate, "cpc_code", ""),
            )

            self._set_cell(
                row.cells[1],
                self._get(candidate, "title", ""),
            )

            self._set_cell(
                row.cells[2],
                self._get(candidate, "feature_id", "") or "",
            )

            self._set_cell(
                row.cells[3],
                self._get(candidate, "relevance", ""),
            )

            self._set_cell(
                row.cells[4],
                self._get(candidate, "reason", ""),
            )

        document.add_paragraph()

    # =========================================================
    # 5 & 6. SEARCH QUERIES
    # =========================================================

    def _add_search_queries(self, document, query_analysis, heading, search_system):

        document.add_heading(heading, level=1)

        if search_system == "USPTO":
            document.add_paragraph(
                "The following search statements are formatted "
                "for the USPTO Patent Public Search Advanced "
                "Search interface. Four independent search paths "
                "are provided for patent titles, descriptions, "
                "claims and CPC classifications. The text searches "
                "use the technical terminology generated from the "
                "invention, while the CPC search uses classifications "
                "identified during CPC analysis."
            )

        elif search_system == "Espacenet":
            document.add_paragraph(
                "The following search statements are formatted "
                "for the EPO Espacenet Smart Search interface. "
                "Four independent search paths are provided for "
                "patent titles, descriptions, claims and CPC "
                "classifications. The text searches use the "
                "technical terminology generated from the invention, "
                "while the CPC search uses classifications identified "
                "during CPC analysis."
            )
        document.add_paragraph(
            "The queries documented below are search statements prepared for execution in the respective patent search system. This report does not represent the results of executing these searches."
        )
        for index, item in enumerate(self._get(query_analysis, "queries", []) or [], start=1):
            # -----------------------------------------------------
            # Query title
            # -----------------------------------------------------

            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{index}. {self._get(item, 'query_name', f'Query {index}')}")
            run.bold = True
            run.font.size = Pt(11)

            # -----------------------------------------------------
            # Priority
            # -----------------------------------------------------

            paragraph = document.add_paragraph()
            run = paragraph.add_run("Priority: ")
            run.bold = True
            priority_run = paragraph.add_run(str(self._get(item, "priority", "Not specified")))

            # -----------------------------------------------------
            # Purpose
            # -----------------------------------------------------

            paragraph = document.add_paragraph()
            run = paragraph.add_run("Purpose: ")
            run.bold = True
            paragraph.add_run(str(self._get(item, "purpose", "Not specified")))

            # -----------------------------------------------------
            # Target Features
            #
            # Kept for backward compatibility if feature_ids
            # exists in either query model.
            # -----------------------------------------------------

            feature_ids = self._get(item, "feature_ids", []) or []
            if feature_ids:
                paragraph = document.add_paragraph()
                run = paragraph.add_run("Target Features: ")
                run.bold = True
                paragraph.add_run(", ".join(str(value) for value in feature_ids))

            # -----------------------------------------------------
            # Query
            # -----------------------------------------------------

            table = document.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.rows[0].cells[0]
            self._shade_cell(cell, "F3F3F3")
            self._set_cell(cell, self._get(item, "query", ""))
            document.add_paragraph()

    # =========================================================
    # 7. NON-PATENT LITERATURE SEARCH RESULTS
    # =========================================================

    def _add_npl_results(self, document, results, summary):
        document.add_heading("7. Non-Patent Literature (NPL) Search Results", level=1)
        document.add_paragraph(
            "This section lists scholarly/preprint records discovered by the automated search. "
            "Relevance scores are machine-assisted triage signals, not legal conclusions. "
            "Source records should be opened and reviewed by the patent practitioner."
        )

        summary = summary if isinstance(summary, dict) else {}
        status = summary.get("status", "not_run")
        metrics = [
            ("Search status", status),
            ("Queries issued", summary.get("query_count", 0)),
            ("Provider requests", summary.get("provider_request_count", 0)),
            ("Raw records", summary.get("raw_result_count", len(results))),
            ("Unique records", summary.get("deduplicated_count", len(results))),
            ("LLM-ranked candidates", summary.get("llm_rerank_count", 0)),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in metrics:
            cells = table.add_row().cells
            self._set_cell(cells[0], label, bold=True)
            self._shade_cell(cells[0], "E7E6E6")
            self._set_cell(cells[1], value if value is not None else "—")

        provider_calls = summary.get("provider_calls", []) or []
        if provider_calls:
            document.add_heading("7.1 Provider Diagnostics", level=2)
            diagnostics = document.add_table(rows=1, cols=5)
            diagnostics.style = "Table Grid"
            for i, heading in enumerate(("Provider", "Category", "Records", "Seconds", "Error")):
                self._set_cell(diagnostics.rows[0].cells[i], heading, bold=True)
                self._shade_cell(diagnostics.rows[0].cells[i], "D9EAF7")
            for call in provider_calls:
                cells = diagnostics.add_row().cells
                for i, key in enumerate(("provider", "category", "count", "elapsed_seconds", "error")):
                    value = call.get(key, "")
                    if key == "elapsed_seconds" and isinstance(value, (int, float)):
                        value = f"{value:.2f}"
                    self._set_cell(cells[i], value if value not in (None, "") else "—")

        if not results:
            document.add_paragraph("No NPL records were returned. This may reflect search coverage, query formulation, provider availability or lack of indexed results; it should not be interpreted as evidence that no relevant literature exists.")
            return

        document.add_heading("7.2 Ranked and Discovered References", level=2)
        for index, record in enumerate(results, start=1):
            title = self._get(record, "title", "Untitled record") or "Untitled record"
            rank = self._get(record, "relevance_rank")
            score = self._get(record, "relevance_score")
            label = f"{rank}. " if rank else f"Record {index}. "
            paragraph = document.add_paragraph()
            run = paragraph.add_run(label + str(title))
            run.bold = True
            if score is not None:
                paragraph.add_run(f"  |  Relevance score: {score}/100")

            authors = self._as_list(self._get(record, "authors", []))
            date = self._get(record, "publication_date") or self._get(record, "year") or "Date not available"
            venue = self._get(record, "venue") or "Venue not available"
            providers = self._as_list(self._get(record, "source_providers", []))
            if not providers:
                providers = self._as_list(self._get(record, "source_provider", []))
            doi = self._get(record, "doi")
            arxiv_id = self._get(record, "arxiv_id")
            url = self._get(record, "url") or self._get(record, "open_access_url")

            metadata = document.add_paragraph()
            metadata.add_run("Authors: ").bold = True
            metadata.add_run(", ".join(str(a) for a in authors[:8]) or "Not available")
            metadata.add_run("  |  Date: ").bold = True
            metadata.add_run(str(date))
            metadata.add_run("  |  Venue: ").bold = True
            metadata.add_run(str(venue))

            identifiers = document.add_paragraph()
            identifiers.add_run("Source(s): ").bold = True
            identifiers.add_run(", ".join(str(p) for p in providers) or "Not specified")
            if doi:
                identifiers.add_run("  |  DOI: ").bold = True
                identifiers.add_run(str(doi))
            if arxiv_id:
                identifiers.add_run("  |  arXiv ID: ").bold = True
                identifiers.add_run(str(arxiv_id))
            if url:
                link_p = document.add_paragraph()
                link_p.add_run("Source link: ").bold = True
                link_p.add_run(str(url))

            abstract = self._get(record, "abstract")
            tldr = self._get(record, "tldr") or self._get(record, "semantic_scholar_tldr")
            if abstract:
                p = document.add_paragraph()
                p.add_run("Abstract: ").bold = True
                p.add_run(str(abstract)[:4000])
            if tldr:
                p = document.add_paragraph()
                p.add_run("Semantic Scholar TLDR: ").bold = True
                p.add_run(str(tldr)[:1200])

            matched = self._as_list(self._get(record, "matched_feature_ids", []))
            if matched:
                p = document.add_paragraph()
                p.add_run("Matched invention features: ").bold = True
                p.add_run(", ".join(str(x) for x in matched))
            rationale = self._get(record, "relevance_rationale")
            if rationale:
                p = document.add_paragraph()
                p.add_run("Relevance rationale: ").bold = True
                p.add_run(str(rationale)[:2000])
            evidence = self._as_list(self._get(record, "evidence_phrases", []))
            if evidence:
                p = document.add_paragraph()
                p.add_run("Evidence phrases: ").bold = True
                p.add_run("; ".join(str(x) for x in evidence[:8]))
            document.add_paragraph()

    # =========================================================
    # 8. METHODOLOGY
    # =========================================================

    def _add_methodology(self, document, npl_summary=None):
        document.add_heading("8. Methodology and Scope", level=1)
        points = [
            "The client-provided IDF was converted into a structured representation of structural, procedural/process and functional features.",
            "Technical keywords, synonyms and alternative terminology were generated independently for each feature category and then consolidated.",
            "CPC candidates were retrieved using semantic similarity against the CPC 2026.08 corpus.",
            "Retrieved CPC candidates were evaluated for technical relevance using the language model.",
            "USPTO and Espacenet search queries were constructed from the generated search vocabulary and CPC classifications using the respective search-system field syntax.",
            "NPL discovery searched available scholarly/preprint metadata sources, normalized records, deduplicated identifiers and used machine-assisted relevance triage when configured.",
            "NPL results are limited by provider coverage, indexing, query formulation and available metadata. Absence of a result does not establish absence of relevant prior art.",
            "The generated queries and NPL references are intended to support professional searching and should be reviewed and refined by the patent practitioner.",
            "This report is not, by itself, a legal opinion on novelty, inventive step or patentability.",
        ]
        for point in points:
            document.add_paragraph(point, style="List Bullet")

    # =========================================================
    # 9. TRACEABILITY
    # =========================================================

    def _add_traceability(self, document, session_id, npl_summary=None):
        document.add_heading("9. Traceability", level=1)
        summary = npl_summary if isinstance(npl_summary, dict) else {}
        stages = (
            "Decomposition → Keywords → CPC Retrieval/Evaluation → USPTO Query Generation → "
            "Espacenet Query Generation → Semantic Scholar/OpenAlex/arXiv NPL Search → "
            "Deduplication and Relevance Triage → Client Report"
        )
        values = [
            ("Session ID", session_id),
            ("Analysis Date", datetime.now().isoformat(timespec="seconds")),
            ("CPC Version", "2026.08"),
            ("Analysis Stages", stages),
            ("NPL Search Status", summary.get("status", "not_run")),
            ("NPL Unique Records", summary.get("deduplicated_count", 0)),
            ("Report Status", "Generated from structured graph state"),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in values:
            cells = table.add_row().cells
            self._set_cell(cells[0], label, bold=True)
            self._shade_cell(cells[0], "E7E6E6")
            self._set_cell(cells[1], value if value is not None else "—")

    # =========================================================
    # CONFIDENTIALITY
    # =========================================================

    def _add_confidentiality_notice(
        self,
        document,
    ):

        document.add_paragraph()

        paragraph = document.add_paragraph()

        run = paragraph.add_run("Confidentiality Notice")

        run.bold = True

        document.add_paragraph(
            "This report is intended for the client and "
            "authorized representatives only. It should "
            "be reviewed by the patent practitioner before "
            "being relied upon for any filing, prosecution, "
            "freedom-to-operate, validity or other legal "
            "decision."
        )

    @staticmethod
    def _get(obj, key, default=None):
        """Read a field from either a dictionary or an object/Pydantic model."""
        if isinstance(obj, dict):
            return obj.get(key, default)
        return getattr(obj, key, default)

    @staticmethod
    def _as_list(value):
        if value is None:
            return []
        if isinstance(value, list):
            return value
        if isinstance(value, tuple):
            return list(value)
        if isinstance(value, str):
            return [value] if value else []
        try:
            return list(value)
        except TypeError:
            return [value]

    # =========================================================
    # TABLE HELPERS
    # =========================================================

    @staticmethod
    def _set_cell(
        cell,
        text,
        bold=False,
    ):

        cell.text = ""

        paragraph = cell.paragraphs[0]

        run = paragraph.add_run(str(text))

        run.bold = bold
        run.font.size = Pt(9)

        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    @staticmethod
    def _shade_cell(
        cell,
        fill,
    ):

        tcPr = cell._tc.get_or_add_tcPr()

        shd = OxmlElement("w:shd")

        shd.set(
            qn("w:fill"),
            fill,
        )

        tcPr.append(shd)
