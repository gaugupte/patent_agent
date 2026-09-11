from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field
from langchain_ollama import ChatOllama

from patentability_rules import (
    SECTION_3_RULES,
    SECTION_4_RULE,
)


# ============================================================
# CONSTANTS
# ============================================================

ALLOWED_STATUSES = {
    "NO APPARENT ISSUE",
    "POTENTIAL ISSUE",
    "LIKELY EXCLUDED",
    "INSUFFICIENT INFORMATION",
    "NOT APPLICABLE",
}


# ============================================================
# INPUT
# ============================================================


def read_input(path: str) -> str:

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"Input file not found: {file_path}")

    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        return file_path.read_text(encoding="utf-8")

    if suffix == ".docx":
        doc = Document(str(file_path))

        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if suffix == ".pdf":
        try:
            import pymupdf
        except ImportError:
            raise RuntimeError("PDF input requires PyMuPDF. Install with: pip install pymupdf")

        pdf = pymupdf.open(str(file_path))

        pages = []

        for page in pdf:
            pages.append(page.get_text())

        return "\n".join(pages)

    raise ValueError("Unsupported input format. Use .txt, .docx or .pdf.")


# ============================================================
# SOURCE BLOCKS
# ============================================================


def build_source_blocks(text: str) -> list[dict]:

    blocks = []

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        blocks.append(
            {
                "id": f"P{len(blocks) + 1:04d}",
                "text": line,
            }
        )

    return blocks


def format_source_blocks(
    blocks: list[dict],
) -> str:

    return "\n".join(f"[{b['id']}] {b['text']}" for b in blocks)


# ============================================================
# MODELS
# ============================================================


class DocumentProfile(BaseModel):
    invention_title: str = ""

    title_evidence: list[str] = Field(default_factory=list)

    document_type: str = ""

    technical_field: str = ""

    invention_summary: str = ""

    technical_features: list[str] = Field(default_factory=list)

    software_or_computing_features: list[str] = Field(default_factory=list)

    business_or_administrative_features: list[str] = Field(default_factory=list)

    medical_or_treatment_features: list[str] = Field(default_factory=list)

    operation_or_method: list[str] = Field(default_factory=list)

    technical_effects: list[str] = Field(default_factory=list)

    has_claims: bool = False

    has_abstract: bool = False

    has_drawings: bool = False


class RuleAssessment(BaseModel):
    rule_id: str

    title: str

    status: str

    explanation: str

    supporting_text: list[str] = Field(default_factory=list)

    issues: list[str] = Field(default_factory=list)


class AssessmentBatch(BaseModel):
    assessments: list[RuleAssessment]


# ============================================================
# LLM
# ============================================================


def create_llm(model_name: str):

    return ChatOllama(
        model=model_name,
        temperature=0,
    )


def invoke_structured(
    llm,
    schema,
    prompt: str,
):

    return llm.with_structured_output(schema).invoke(prompt)


# ============================================================
# DOCUMENT PROFILE
# ============================================================


def extract_document_profile(
    llm,
    source_blocks: list[dict],
) -> DocumentProfile:

    source = format_source_blocks(source_blocks)

    prompt = f"""
You are extracting factual information from an Indian patent document.

IMPORTANT:
You are NOT performing patentability analysis.

Use ONLY information actually present in the supplied source.

DO NOT:
- invent a title
- invent technical features
- infer claims
- infer an abstract
- infer technical effects that are not described
- assess novelty
- assess inventive step
- assess prior art

TITLE:
Only return an invention_title if an actual title or heading
appears in the source.

title_evidence must contain the source block ID where the title
appears.

If no actual title can be identified:
invention_title = ""
title_evidence = []

DOCUMENT TYPE:
Determine whether the document appears to be:
- provisional specification
- complete specification
- invention disclosure
- patent application
- unknown

Do not assume legal filing status merely from the filename.

SOURCE:

{source}
"""

    return invoke_structured(
        llm,
        DocumentProfile,
        prompt,
    )


# ============================================================
# SECTION 3 / 4
# ============================================================


def evaluate_rules(
    llm,
    rules: list[dict],
    profile: DocumentProfile,
    source_blocks: list[dict],
    category: str,
) -> list[RuleAssessment]:

    source = format_source_blocks(source_blocks)

    rules_text = "\n\n".join(
        f"""
RULE ID: {r["id"]}
TITLE: {r["title"]}
TEST: {r.get("test", "")}
"""
        for r in rules
    )

    profile_text = profile.model_dump_json(indent=2)

    prompt = f"""
You are conducting a PRELIMINARY statutory screening
under the Indian Patents Act.

CATEGORY:

{category}

Evaluate ONLY the rules supplied below.

Do NOT introduce any other patent-law categories.

In particular, do NOT evaluate:
- novelty
- anticipation
- prior art
- inventive step
- obviousness
- patentability in general
- manner of manufacture
- usefulness
- foreign patent law

For every rule:

NO APPARENT ISSUE
= the supplied information does not reveal an apparent
statutory problem.

POTENTIAL ISSUE
= there is a meaningful concern, but the available
information does not establish the exclusion conclusively.

LIKELY EXCLUDED
= the supplied information strongly indicates that
the statutory exclusion applies.

INSUFFICIENT INFORMATION
= the source does not contain enough information.

NOT APPLICABLE
= the subject matter clearly falls outside the rule.

CRITICAL EVIDENCE RULE:

Do not merely say "the invention is not X."

Explain the relevant subject matter and use source
references such as P0004 whenever possible.

Never fabricate a source reference.

DOCUMENT PROFILE:

{profile_text}

SOURCE:

{source}

RULES:

{rules_text}
"""

    result = invoke_structured(
        llm,
        AssessmentBatch,
        prompt,
    )

    valid_ids = {r["id"] for r in rules}

    assessments = []

    for assessment in result.assessments:
        if assessment.rule_id not in valid_ids:
            continue

        if assessment.status not in ALLOWED_STATUSES:
            assessment.status = "INSUFFICIENT INFORMATION"

        assessments.append(assessment)

    existing = {a.rule_id for a in assessments}

    for rule in rules:
        if rule["id"] in existing:
            continue

        assessments.append(
            RuleAssessment(
                rule_id=rule["id"],
                title=rule["title"],
                status="INSUFFICIENT INFORMATION",
                explanation=("No valid assessment was returned for this statutory rule."),
            )
        )

    order = {rule["id"]: i for i, rule in enumerate(rules)}

    assessments.sort(
        key=lambda x: order.get(
            x.rule_id,
            999,
        )
    )

    return assessments


# ============================================================
# SECTION 3(k) — SPECIAL ANALYSIS
# ============================================================


class Section3KAnalysis(BaseModel):
    technical_character: list[str] = Field(default_factory=list)

    business_or_administrative_character: list[str] = Field(default_factory=list)

    computer_implementation: list[str] = Field(default_factory=list)

    technical_effects: list[str] = Field(default_factory=list)

    evidence: list[str] = Field(default_factory=list)

    conclusion: str = ""


def analyze_section_3k(
    llm,
    profile: DocumentProfile,
    source_blocks: list[dict],
) -> Section3KAnalysis:

    source = format_source_blocks(source_blocks)

    prompt = f"""
Analyse ONLY the Section 3(k) subject matter of the
supplied Indian patent document.

Do not decide novelty or inventive step.

The purpose is to distinguish:

1. technical subject matter;
2. business or administrative subject matter;
3. computer/software implementation;
4. disclosed technical effects.

Use only facts actually disclosed.

Do NOT treat the mere presence of:
- a computer,
- processor,
- database,
- mobile application,
- internet,
- software,
- data processing

as automatically establishing a technical effect.

Conversely, do not assume that software involvement automatically
means Section 3(k) exclusion.

Identify the actual technical contribution described by the document.

Provide source references such as P0004.

DOCUMENT PROFILE:

{profile.model_dump_json(indent=2)}

SOURCE:

{source}
"""

    return invoke_structured(
        llm,
        Section3KAnalysis,
        prompt,
    )


# ============================================================
# SECTION 3(k) ASSESSMENT
# ============================================================


def build_section_3k_assessment(
    analysis: Section3KAnalysis,
) -> RuleAssessment:

    if analysis.conclusion.strip():
        conclusion = analysis.conclusion.lower()

        if "likely excluded" in conclusion or "excluded under 3(k)" in conclusion:
            status = "LIKELY EXCLUDED"

        elif "potential issue" in conclusion or "possible exclusion" in conclusion:
            status = "POTENTIAL ISSUE"

        elif "insufficient" in conclusion:
            status = "INSUFFICIENT INFORMATION"

        else:
            status = "NO APPARENT ISSUE"

    else:
        if analysis.business_or_administrative_character and not analysis.technical_effects:
            status = "POTENTIAL ISSUE"

        elif analysis.technical_effects:
            status = "NO APPARENT ISSUE"

        else:
            status = "INSUFFICIENT INFORMATION"

    explanation_parts = []

    if analysis.technical_character:
        explanation_parts.append("Technical character identified: " + "; ".join(analysis.technical_character))

    if analysis.business_or_administrative_character:
        explanation_parts.append("Business/administrative character identified: " + "; ".join(analysis.business_or_administrative_character))

    if analysis.computer_implementation:
        explanation_parts.append("Computer implementation: " + "; ".join(analysis.computer_implementation))

    if analysis.technical_effects:
        explanation_parts.append("Technical effects identified: " + "; ".join(analysis.technical_effects))

    if analysis.conclusion:
        explanation_parts.append("Conclusion: " + analysis.conclusion)

    return RuleAssessment(
        rule_id="3(k)",
        title=("Mathematical method, business method, computer programme per se or algorithms"),
        status=status,
        explanation=" ".join(explanation_parts),
        supporting_text=analysis.evidence,
    )


# ============================================================
# TEXTUAL HEADINGS
# ============================================================


def normalize_heading(text: str) -> str:

    return re.sub(
        r"[^A-Z0-9 ]",
        "",
        text.upper(),
    ).strip()


def find_heading_position(
    text: str,
    patterns: list[str],
) -> int | None:

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE | re.MULTILINE,
        )

        if match:
            return match.start()

    return None


# ============================================================
# ABSTRACT EXTRACTION
# ============================================================


def extract_abstract(
    text: str,
) -> str:

    # Require ABSTRACT to look like an actual heading.
    start_pattern = re.compile(
        r"(?im)^[ \t]*(?:ABSTRACT|"
        r"ABSTRACT OF THE INVENTION|"
        r"ABSTRACT OF THE DISCLOSURE)"
        r"[ \t]*[:\-]?[ \t]*$"
    )

    match = start_pattern.search(text)

    if not match:
        return ""

    start = match.end()

    end_patterns = [
        r"(?im)^[ \t]*CLAIMS[ \t]*$",
        r"(?im)^[ \t]*FIELD OF INVENTION[ \t]*$",
        r"(?im)^[ \t]*BACKGROUND[ \t]*$",
        r"(?im)^[ \t]*BACKGROUND OF THE INVENTION[ \t]*$",
        r"(?im)^[ \t]*SUMMARY[ \t]*$",
        r"(?im)^[ \t]*SUMMARY OF THE INVENTION[ \t]*$",
        r"(?im)^[ \t]*DETAILED DESCRIPTION[ \t]*$",
        r"(?im)^[ \t]*DESCRIPTION[ \t]*$",
    ]

    end_positions = []

    for pattern in end_patterns:
        m = re.search(
            pattern,
            text[start:],
        )

        if m:
            end_positions.append(start + m.start())

    if end_positions:
        end = min(end_positions)
    else:
        # Do NOT consume the entire document.
        # A real abstract is normally relatively short.
        candidate = text[start:]

        paragraphs = [
            p.strip()
            for p in re.split(
                r"\n\s*\n",
                candidate,
            )
            if p.strip()
        ]

        candidate = "\n\n".join(paragraphs[:3])

        words = candidate.split()

        if len(words) > 250:
            return ""

        return candidate.strip()

    abstract = text[start:end].strip()

    words = abstract.split()

    # If extraction looks suspiciously large,
    # reject it rather than reporting a fake abstract.
    if len(words) > 250:
        return ""

    return abstract


# ============================================================
# ABSTRACT CHECKS
# ============================================================


def abstract_checks(
    profile: DocumentProfile,
    source_text: str,
) -> list[RuleAssessment]:

    abstract = extract_abstract(source_text)

    if not abstract:
        return [
            RuleAssessment(
                rule_id="13(7)(a)",
                title="Abstract begins with title",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable abstract section could be identified in the supplied document."),
            ),
            RuleAssessment(
                rule_id="13(7)(b)",
                title="Contents of abstract",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable abstract section could be identified in the supplied document."),
            ),
            RuleAssessment(
                rule_id="13(7)(c)",
                title="Abstract length",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable abstract section could be identified, so its word count cannot be reliably assessed."),
            ),
            RuleAssessment(
                rule_id="13(7)(d)",
                title="Reference signs in drawings",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable abstract section could be identified."),
            ),
            RuleAssessment(
                rule_id="13(7)(e)",
                title="Abstract as search instrument",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable abstract section could be identified."),
            ),
        ]

    words = re.findall(
        r"\b[\w'-]+\b",
        abstract,
    )

    word_count = len(words)

    title = profile.invention_title.strip()

    title_words = len(title.split()) if title else 0

    # 13(7)(a)
    if title and title_words <= 15:
        status_a = "NO APPARENT ISSUE"
        explanation_a = f"The identified title contains approximately {title_words} words. The abstract section was identified separately in the document."

    elif title and title_words > 15:
        status_a = "POTENTIAL ISSUE"
        explanation_a = f"The identified title contains approximately {title_words} words, exceeding the normal fifteen-word limit."

    else:
        status_a = "INSUFFICIENT INFORMATION"
        explanation_a = "A reliable invention title was not identified, so the abstract/title relationship cannot be reliably assessed."

    return [
        RuleAssessment(
            rule_id="13(7)(a)",
            title="Abstract begins with title",
            status=status_a,
            explanation=explanation_a,
        ),
        RuleAssessment(
            rule_id="13(7)(b)",
            title="Contents of abstract",
            status="INSUFFICIENT INFORMATION",
            explanation=("An abstract is present, but substantive assessment of its technical field, technical advancement and principal use requires textual review."),
        ),
        RuleAssessment(
            rule_id="13(7)(c)",
            title="Abstract length",
            status=("NO APPARENT ISSUE" if word_count <= 150 else "POTENTIAL ISSUE"),
            explanation=(f"The identified abstract contains approximately {word_count} words."),
        ),
        RuleAssessment(
            rule_id="13(7)(d)",
            title="Reference signs in drawings",
            status="INSUFFICIENT INFORMATION",
            explanation=("Comparison of abstract reference signs with drawing reference signs requires examination of the drawings."),
        ),
        RuleAssessment(
            rule_id="13(7)(e)",
            title="Abstract as search instrument",
            status="INSUFFICIENT INFORMATION",
            explanation=("The abstract is present, but suitability as an efficient search instrument requires substantive review."),
        ),
    ]


# ============================================================
# SECTION 10 / RULE 13
# ============================================================


def deterministic_specification_checks(
    profile: DocumentProfile,
    source_text: str,
) -> tuple[
    list[RuleAssessment],
    list[RuleAssessment],
]:

    section_10 = []

    # --------------------------------------------------------
    # 10(1)
    # --------------------------------------------------------

    if profile.invention_title and profile.title_evidence:
        section_10.append(
            RuleAssessment(
                rule_id="10(1)",
                title="Specification must begin with a title",
                status="NO APPARENT ISSUE",
                explanation=("An invention title was identified in the source document."),
                supporting_text=profile.title_evidence,
            )
        )

    else:
        section_10.append(
            RuleAssessment(
                rule_id="10(1)",
                title="Specification must begin with a title",
                status="INSUFFICIENT INFORMATION",
                explanation=("No reliable invention title could be identified from the supplied document. This is not treated as a confirmed legal defect by the analyzer."),
            )
        )

    # --------------------------------------------------------
    # 10(4)(a)
    # --------------------------------------------------------

    if profile.technical_features or profile.operation_or_method:
        status = "NO APPARENT ISSUE"

        explanation = (
            "The document contains identifiable technical features and/or an operation or method. This screen does not determine whether the disclosure is legally sufficient in every respect."
        )

    else:
        status = "INSUFFICIENT INFORMATION"

        explanation = "Insufficient identifiable technical description was extracted to perform this screening."

    section_10.append(
        RuleAssessment(
            rule_id="10(4)(a)",
            title="Full and particular description",
            status=status,
            explanation=explanation,
        )
    )

    # --------------------------------------------------------
    # 10(4)(b)
    # --------------------------------------------------------

    section_10.append(
        RuleAssessment(
            rule_id="10(4)(b)",
            title="Best method of performing the invention",
            status="INSUFFICIENT INFORMATION",
            explanation=(
                "The supplied document does not provide a "
                "reliable basis for determining the best method "
                "known to the applicant. This normally requires "
                "substantive review of the complete specification "
                "and the applicant's knowledge."
            ),
        )
    )

    # --------------------------------------------------------
    # 10(4)(c)
    # --------------------------------------------------------

    section_10.append(
        RuleAssessment(
            rule_id="10(4)(c)",
            title="Claims defining scope",
            status=("NO APPARENT ISSUE" if profile.has_claims else "INSUFFICIENT INFORMATION"),
            explanation=(
                "Claims appear to be present."
                if profile.has_claims
                else "No claims were identified. If the document is a provisional specification or invention disclosure, absence of claims should not automatically be treated as a defect."
            ),
        )
    )

    # --------------------------------------------------------
    # 10(4)(d)
    # --------------------------------------------------------

    section_10.append(
        RuleAssessment(
            rule_id="10(4)(d)",
            title="Abstract",
            status=("NO APPARENT ISSUE" if profile.has_abstract else "INSUFFICIENT INFORMATION"),
            explanation=("An abstract appears to be present." if profile.has_abstract else "No reliable abstract was identified in the supplied document."),
        )
    )

    # --------------------------------------------------------
    # 10(5)
    # --------------------------------------------------------

    section_10.append(
        RuleAssessment(
            rule_id="10(5)",
            title="Claims must be clear, succinct and fairly based",
            status="INSUFFICIENT INFORMATION",
            explanation=("A substantive assessment of claim clarity, succinctness, fair basis and single inventive concept requires examination of the actual claims."),
        )
    )

    # ========================================================
    # RULE 13
    # ========================================================

    rule_13 = []

    rule_13.append(
        RuleAssessment(
            rule_id="13(1)",
            title="Specification in prescribed form",
            status="INSUFFICIENT INFORMATION",
            explanation=("Form compliance cannot reliably be established from extracted document text alone."),
        )
    )

    rule_13.append(
        RuleAssessment(
            rule_id="13(4)",
            title="Drawings where necessary",
            status=("NO APPARENT ISSUE" if profile.has_drawings else "INSUFFICIENT INFORMATION"),
            explanation=("Drawings appear to be present." if profile.has_drawings else "No drawings were identified. The analyzer cannot determine from text alone whether drawings are necessary."),
        )
    )

    rule_13.extend(
        abstract_checks(
            profile,
            source_text,
        )
    )

    return section_10, rule_13


# ============================================================
# REPORT
# ============================================================


def build_report(
    profile: DocumentProfile,
    section_3: list[RuleAssessment],
    section_4: list[RuleAssessment],
    section_10: list[RuleAssessment],
    rule_13: list[RuleAssessment],
) -> str:

    all_assessments = section_3 + section_4 + section_10 + rule_13

    if any(a.status == "LIKELY EXCLUDED" for a in all_assessments):
        overall = "POTENTIAL STATUTORY EXCLUSION IDENTIFIED"

    elif any(a.status == "POTENTIAL ISSUE" for a in all_assessments):
        overall = "POTENTIAL STATUTORY / SPECIFICATION ISSUE IDENTIFIED"

    elif any(a.status == "INSUFFICIENT INFORMATION" for a in all_assessments):
        overall = "NO APPARENT ISSUE — SOME ITEMS REQUIRE FURTHER REVIEW"

    else:
        overall = "NO APPARENT ISSUE"

    lines = []

    lines.append("INDIAN PATENTABILITY SCREENING REPORT")

    lines.append("")

    lines.append("INVENTION TITLE")

    lines.append(profile.invention_title or "Not reliably identified")

    lines.append("")

    lines.append("DOCUMENT TYPE")

    lines.append(profile.document_type or "Unknown")

    lines.append("")

    lines.append("SCOPE OF SCREENING")

    lines.append("This report is a preliminary statutory screening based only on the supplied document.")

    lines.append("It assesses statutory exclusions under Sections 3 and 4 and selected specification requirements under Section 10 and Rule 13.")

    lines.append("It does NOT assess novelty, anticipation, prior art, inventive step or obviousness.")

    lines.append("")

    lines.append("OVERALL SCREENING")

    lines.append(overall)

    # --------------------------------------------------------
    # Section 3
    # --------------------------------------------------------

    lines.append("")
    lines.append("SECTION 3 — NON-PATENTABLE INVENTIONS")

    for assessment in section_3:
        add_assessment(
            lines,
            assessment,
        )

    # --------------------------------------------------------
    # Section 4
    # --------------------------------------------------------

    lines.append("")
    lines.append("SECTION 4 — ATOMIC ENERGY")

    for assessment in section_4:
        add_assessment(
            lines,
            assessment,
        )

    # --------------------------------------------------------
    # Section 10
    # --------------------------------------------------------

    lines.append("")
    lines.append("SECTION 10 — CONTENTS OF SPECIFICATIONS")

    for assessment in section_10:
        add_assessment(
            lines,
            assessment,
        )

    # --------------------------------------------------------
    # Rule 13
    # --------------------------------------------------------

    lines.append("")
    lines.append("RULE 13 — SPECIFICATION / ABSTRACT REQUIREMENTS")

    for assessment in rule_13:
        add_assessment(
            lines,
            assessment,
        )

    # --------------------------------------------------------
    # Key issues
    # --------------------------------------------------------

    lines.append("")
    lines.append("KEY ISSUES")

    issues_found = [
        a
        for a in all_assessments
        if a.status
        in {
            "POTENTIAL ISSUE",
            "LIKELY EXCLUDED",
        }
    ]

    if issues_found:
        for issue in issues_found:
            lines.append(f"- {issue.rule_id}: {issue.explanation}")

    else:
        lines.append("- No potential statutory exclusion was identified.")

    # --------------------------------------------------------
    # Limitations
    # --------------------------------------------------------

    lines.append("")
    lines.append("IMPORTANT LIMITATIONS")

    lines.append("This is a preliminary screening and not a legal opinion.")

    lines.append("The analyzer does not determine novelty, prior art, inventive step or obviousness.")

    return "\n".join(lines)


def add_assessment(
    lines: list[str],
    assessment: RuleAssessment,
):

    lines.append("")
    lines.append(f"{assessment.rule_id} — {assessment.title}")

    lines.append(f"STATUS: {assessment.status}")

    lines.append(f"Explanation: {assessment.explanation}")

    if assessment.supporting_text:
        lines.append("Supporting source references:")

        for source in assessment.supporting_text:
            lines.append(f"  - {source}")

    if assessment.issues:
        lines.append("Issues:")

        for issue in assessment.issues:
            lines.append(f"  - {issue}")


# ============================================================
# DOCX
# ============================================================


def save_docx(
    text: str,
    output_path: str,
):

    doc = Document()

    for line in text.splitlines():
        if not line.strip():
            doc.add_paragraph()
            continue

        if line.startswith(
            (
                "INDIAN PATENTABILITY",
                "INVENTION TITLE",
                "DOCUMENT TYPE",
                "SCOPE OF SCREENING",
                "OVERALL SCREENING",
                "SECTION 3",
                "SECTION 4",
                "SECTION 10",
                "RULE 13",
                "KEY ISSUES",
                "IMPORTANT LIMITATIONS",
            )
        ):
            doc.add_heading(
                line,
                level=1,
            )

        elif re.match(
            r"^(3\([a-z]\)|4|10\([^)]+\)|13\([^)]+\))\s+—",
            line,
        ):
            doc.add_heading(
                line,
                level=2,
            )

        elif line.startswith("STATUS:"):
            p = doc.add_paragraph()

            p.add_run(line).bold = True

        else:
            doc.add_paragraph(line)

    doc.save(output_path)


# ============================================================
# MAIN
# ============================================================


def main():

    parser = argparse.ArgumentParser(description=("Indian Patentability Statutory Screening"))

    parser.add_argument(
        "input_file",
        help=("TXT, DOCX or PDF invention/specification file"),
    )

    parser.add_argument(
        "-m",
        "--model",
        required=True,
        help=("Ollama model, e.g. llama3.1:8b"),
    )

    parser.add_argument(
        "-o",
        "--output",
        required=True,
        help=("Output DOCX file"),
    )

    args = parser.parse_args()

    print("Reading input...")

    source_text = read_input(args.input_file)

    if not source_text.strip():
        raise RuntimeError("Input document contains no extractable text.")

    source_blocks = build_source_blocks(source_text)

    print(f"Extracted {len(source_blocks)} source blocks.")

    print(f"Starting Ollama model: {args.model}")

    llm = create_llm(args.model)

    # --------------------------------------------------------
    # Stage 1
    # --------------------------------------------------------

    print("Extracting document profile...")

    profile = extract_document_profile(
        llm,
        source_blocks,
    )

    # --------------------------------------------------------
    # Section 3
    # --------------------------------------------------------

    print("Evaluating Section 3...")

    section_3_rules = [r for r in SECTION_3_RULES if r["id"] != "3(k)"]

    section_3 = evaluate_rules(
        llm,
        section_3_rules,
        profile,
        source_blocks,
        "SECTION 3 — NON-PATENTABLE INVENTIONS",
    )

    # --------------------------------------------------------
    # Section 3(k)
    # --------------------------------------------------------

    print("Performing dedicated Section 3(k) analysis...")

    section_3k_analysis = analyze_section_3k(
        llm,
        profile,
        source_blocks,
    )

    section_3.append(build_section_3k_assessment(section_3k_analysis))

    # Preserve statutory order
    section_3_order = {rule["id"]: i for i, rule in enumerate(SECTION_3_RULES)}

    section_3.sort(
        key=lambda x: section_3_order.get(
            x.rule_id,
            999,
        )
    )

    # --------------------------------------------------------
    # Section 4
    # --------------------------------------------------------

    print("Evaluating Section 4...")

    section_4 = evaluate_rules(
        llm,
        [SECTION_4_RULE],
        profile,
        source_blocks,
        "SECTION 4 — ATOMIC ENERGY",
    )

    # --------------------------------------------------------
    # Section 10 / Rule 13
    # --------------------------------------------------------

    print("Checking Section 10 and Rule 13...")

    section_10, rule_13 = deterministic_specification_checks(
        profile,
        source_text,
    )

    # --------------------------------------------------------
    # Report
    # --------------------------------------------------------

    report = build_report(
        profile,
        section_3,
        section_4,
        section_10,
        rule_13,
    )

    save_docx(
        report,
        args.output,
    )

    print("")
    print("========================================")
    print("PATENTABILITY SCREENING COMPLETE")
    print("========================================")

    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
